#!/usr/bin/env python3
"""Envoie les 33 courriels anglophones, un aux 30 secondes.

Tourne en parallele de envoi_fr.py, donc il n'ecrit JAMAIS dans le classeur :
les deux processus se marcheraient dessus et le .xlsx y laisserait sa peau.
Seul `envoi_en_journal.json` est ecrit, apres chaque envoi, ce qui rend la
reprise possible. La colonne « Envoyé le » de la feuille anglaise se remplit
apres coup, quand les deux envois sont termines.

Ces brouillons portent deja la signature anglaise dans leur corps, et ils ne
recoivent pas la presentation d'entreprise : elle est en francais, et un
pupitre anglophone qui recoit un PDF illisible apprend surtout qu'on ne l'a
pas regarde.

    python3 envoi_en.py                 simulation
    python3 envoi_en.py --envoyer       envoie pour de vrai
"""
import base64
import datetime
import json
import os
import subprocess
import sys
import time

import openpyxl

from voix_gabriel import en_html

ICI = os.path.dirname(os.path.abspath(__file__))
CLIENT = os.path.join(os.path.dirname(ICI), "missive_client.js")
CLASSEUR = os.path.join(ICI, "Lasclay_v2.xlsx")
JOURNAL = os.path.join(ICI, "envoi_en_journal.json")
JOURNAL_FR = os.path.join(ICI, "envoi_fr_journal.json")
PHOTO = os.path.join(ICI, "dragons-plateau.jpg")
DE = "media@lasclay.com"


def appel(charge):
    p = subprocess.run(["node", CLIENT, "send"], input=json.dumps(charge),
                       capture_output=True, text=True, timeout=300)
    if p.returncode != 0:
        return {"erreur": (p.stderr or p.stdout).strip()[:300]}
    try:
        return json.loads(p.stdout)
    except ValueError:
        return {"erreur": p.stdout.strip()[:300]}


def file_attente():
    ws = openpyxl.load_workbook(CLASSEUR)["Presse anglophone"]
    h = {c.value: i for i, c in enumerate(ws[1])}
    file, vues = [], set()
    for r in ws.iter_rows(min_row=2, values_only=True):
        a, b = r[h["Courriel"]], r[h["Brouillon"]]
        if not a or not b or b.startswith("NE PAS"):
            continue
        cle = a.strip().lower()
        if cle in vues:
            continue
        vues.add(cle)
        file.append({"courriel": a, "objet": r[h["Objet"]], "corps": b,
                     "nom": r[h["Média"]]})
    return file


def main(argv):
    envoyer = "--envoyer" in argv
    pause = 30.0
    if "--pause" in argv:
        pause = float(argv[argv.index("--pause") + 1])
    journal = json.load(open(JOURNAL, encoding="utf-8")) if os.path.exists(JOURNAL) else {}
    # Personne ne recoit la version francaise ET l'anglaise.
    deja = set(journal)
    if os.path.exists(JOURNAL_FR):
        deja |= set(json.load(open(JOURNAL_FR, encoding="utf-8")))

    file = file_attente()
    reste = [t for t in file if t["courriel"].strip().lower() not in deja]
    print(f"{len(file)} courriels anglais, {len(file) - len(reste)} déjà servis, "
          f"{len(reste)} à envoyer, un aux {pause:.0f} s "
          f"→ environ {len(reste) * pause / 60:.0f} minutes.", flush=True)
    if not envoyer:
        for t in reste[:5]:
            print(f"  · {t['courriel']:42s} {t['objet']}")
        print("\nSimulation. Relancer avec --envoyer.")
        return 0

    pieces = [{"base64_data": base64.b64encode(open(PHOTO, "rb").read()).decode(),
               "filename": "lasclay-dragons-den.jpg"}]
    faits, echecs = 0, []
    for i, t in enumerate(reste, 1):
        res = appel({"from": DE, "to": [t["courriel"]], "subject": t["objet"],
                     "body": en_html(t["corps"]), "send": True, "attachments": pieces})
        horo = datetime.datetime.now().isoformat(timespec="seconds")
        if res.get("ok"):
            faits += 1
            journal[t["courriel"].strip().lower()] = {
                "quand": horo, "conversation": res.get("conversation")}
            json.dump(journal, open(JOURNAL, "w"), ensure_ascii=False, indent=1)
            print(f"[{i}/{len(reste)}] {horo[11:]} ✓ {t['courriel']:42s} {t['nom']}",
                  flush=True)
        else:
            echecs.append((t["courriel"], res.get("erreur") or res))
            print(f"[{i}/{len(reste)}] {horo[11:]} ✗ {t['courriel']:42s} "
                  f"{res.get('erreur') or res}", flush=True)
        if i < len(reste):
            time.sleep(pause)

    print(f"\n{faits} envoyés, {len(echecs)} en échec.", flush=True)
    for a, e in echecs:
        print(f"  {a} : {e}", flush=True)
    return 1 if echecs else 0


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))
