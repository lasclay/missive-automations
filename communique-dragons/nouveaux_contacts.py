#!/usr/bin/env python3
"""Les contacts trouves le 15 septembre 2026, deux jours avant la diffusion.

Gabriel : « ne pas ecrire a qui que ce soit de CBC. Mais Radio-Canada oui,
incluant emissions de radio. » CBC diffuse l'emission : passer par-dessus son
equipe de publicite pour parler a ses salles de nouvelles se retourne contre
nous. Toutes les adresses @cbc.ca sont donc retirees, y compris celle du
Media Centre que j'avais mise la veille.

Meme regle que pour le reste du dossier : chaque adresse est PUBLIEE par le
media, et la source est notee a cote. Rien n'est deduit d'un patron
« prenom.nom@ » — sauf que, pour La Semaine verte, c'est le media lui-meme qui
publie la liste nominative de son equipe, roles compris. On n'a pas devine, on
a lu.

Le tri par role compte autant que l'adresse : ecrire aux dix-huit personnes de
l'equipe, montage et camera compris, c'est du publipostage. Ne restent que la
redactrice en chef, l'animatrice et les journalistes.
"""
import sys

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill

from voix_gabriel import BENEFICE, media_kit

DRIVE = "https://drive.google.com/drive/folders/1pyCUbfHYQhpXXl4FoCC2RCFXKRvGS5Zr"

# --- francais ---------------------------------------------------------------

QUI = ("Je m'appelle Gabriel Gouveia, fondateur de Lasclay. On isole des vêtements d'hiver avec "
       "une mauvaise herbe, l'asclépiade, qu'on cultive pour sauvegarder un pollinisateur "
       "emblématique et menacé : le papillon monarque.")
MATIERE = ("Ses gousses sont remplies d'une soie creuse, très légère et naturellement hydrophobe. "
           "On la transforme en isolant à Québec, et on en fait des manteaux, des mitaines, des "
           "tuques et des sacs isothermes vendus au Canada et aux États-Unis.")
ANNONCE = ("Ce jeudi 17 septembre, je présente Lasclay dans le premier épisode de la 21e saison "
           "de Dragons' Den, sur CBC et CBC Gem, à 20 h.")
CLOTURE = "Au plaisir et n'hésitez pas à me contacter si vous avez des questions."
SIGNATURE = "Chaleureusement,\n__\nGabriel Gouveia\nCo-fondateur\n+1 (581) 982-5857\nLasclay.com"

AGRICOLE = (
    "Je vous écris parce que l'asclépiade est d'abord une histoire agricole. La filière s'est "
    "cassée en 2018 quand l'usine de Saint-Tite a fermé, et la question n'a pas changé depuis : "
    "y a-t-il un acheteur stable au bout du champ. Nous, on achète encore. Sabin Tremblay, à "
    "L'Ascension-de-Notre-Seigneur au Lac-Saint-Jean, cultive pour nous depuis nos tout débuts, "
    "et je vais l'aider à récolter le 4 octobre — s'il y a un reportage à faire, il est là autant "
    "qu'ici.")

def fr(pourquoi, offre=None):
    bl = [QUI, MATIERE, pourquoi, ANNONCE,
          "Aller présenter notre entreprise et sa mission à la télévision nationale est une "
          "opportunité qui arrive bien rarement. Je voulais vous en faire part et qui sait, "
          "peut-être vous inspirer un sujet." + (" " + offre if offre else ""),
          media_kit("vous"), BENEFICE, CLOTURE, SIGNATURE]
    return "\n\n".join(bl)


DISPO = "Je suis disponible pour une entrevue d'ici jeudi, ou le vendredi 18 au matin."

SEMAINE_VERTE = {
 "johanne.j.lapierre@radio-canada.ca": ("Johanne Lapierre", "Rédactrice en chef"),
 "catherine.mercier@radio-canada.ca": ("Catherine Mercier", "Animatrice"),
 "france.beaudoin@radio-canada.ca": ("France Beaudoin", "Journaliste"),
 "gilbert.begin@radio-canada.ca": ("Gilbert Bégin", "Journaliste"),
 "carine.monat@radio-canada.ca": ("Carine Monat", "Journaliste"),
 "maxime.poire@radio-canada.ca": ("Maxime Poiré", "Journaliste"),
 "marie-maude.pontbriand@radio-canada.ca": ("Marie Maude Pontbriand", "Journaliste"),
 "julie.vaillancourt@radio-canada.ca": ("Julie Vaillancourt", "Journaliste"),
 "marc-yvan.hebert@radio-canada.ca": ("Marc-Yvan Hébert", "Journaliste-réalisateur"),
 "benoit.livernoche@radio-canada.ca": ("Benoit Livernoche", "Journaliste-réalisateur"),
}

TVA = {
 "nouvelles@tva.ca": ("TVA Nouvelles", "Pupitre national", None),
 "nouvelles.quebec@tva.ca": ("TVA Québec", "Pupitre de Québec",
    "Notre atelier est à Québec, et c'est là que la soie d'asclépiade devient de l'isolant."),
 "nouvelles.cjpm@tva.ca": ("TVA Saguenay", "Pupitre du Saguenay",
    "Notre fournisseur du Lac-Saint-Jean, Sabin Tremblay de L'Ascension-de-Notre-Seigneur, "
    "cultive pour nous depuis nos tout débuts."),
 "nouvelles.chem@tva.ca": ("TVA Trois-Rivières", "Pupitre de la Mauricie",
    "L'asclépiade a eu son grand moment industriel chez vous : l'usine de Saint-Tite achetait "
    "90 % des récoltes du Québec avant que la filière se casse en 2018."),
 "nouvelles.sherbrooke@tva.ca": ("TVA Sherbrooke", "Pupitre de l'Estrie",
    "Une des premières usines de transformation de la fibre était à Granby, et l'Estrie compte "
    "encore des producteurs d'asclépiade."),
 "nouvelles.cfer@tva.ca": ("TVA Est-du-Québec", "Pupitre de l'Est-du-Québec", None),
 "quebec@quebecormedia.com": ("Québecor Québec", "Salle de nouvelles", None),
 "montreal@quebecormedia.com": ("Québecor Montréal", "Salle de nouvelles", None),
 "sherbrooke@quebecormedia.com": ("Québecor Sherbrooke", "Salle de nouvelles", None),
 "trois-rivieres@quebecormedia.com": ("Québecor Trois-Rivières", "Salle de nouvelles", None),
 "estquebec@quebecormedia.com": ("Québecor Est-du-Québec", "Salle de nouvelles", None),
}

AFFAIRES = (
    "Je vous écris parce que c'est d'abord une histoire d'affaires : bâtir ses propres procédés "
    "faute de sous-traitant prêt à toucher à la fibre, puis changer de modèle manufacturier l'an "
    "dernier pour tenir, et se retrouver à pitcher à la télévision nationale. La vidéo sur le "
    "changement de modèle est ici : https://www.youtube.com/watch?v=GKyHh-Ok9JU")

# --- anglais ----------------------------------------------------------------

EN_QUI = ("My name is Gabriel Gouveia, founder of Lasclay. We insulate winter clothing with a "
          "weed — milkweed — that we grow to help save an emblematic, threatened pollinator: "
          "the monarch butterfly.")
EN_MATIERE = ("Its pods are filled with a hollow floss, very light and naturally water-repellent. "
              "We turn it into insulation in Quebec City and put it into coats, mitts, toques and "
              "cooler bags sold across Canada and the United States.")
EN_ANNONCE = ("This Thursday, September 17, Lasclay is in the first episode of season 21 of "
              "Dragons' Den, on CBC and CBC Gem at 8 p.m. (8:30 NT).")
EN_OFFRE = (f"If you'd like to cover it — or a colleague might — our media kit is here, with "
            f"images from the Dragons' Den floor and of the company: {DRIVE}")
EN_BENEFICE = ("I'm genuinely hoping this visibility moves sales, which more broadly is very good "
               "news for the Quebec milkweed growers we keep buying from, and for the threatened "
               "monarchs that keep breeding in their fields.")
EN_CLOTURE = "Happy to answer any questions, and available for an interview this week."
EN_SIGNATURE = "Warmly,\n__\nGabriel Gouveia\nCo-founder\n+1 (581) 982-5857\nLasclay.com"

def en(pourquoi):
    return "\n\n".join(["Hello,", EN_QUI, EN_MATIERE, pourquoi, EN_ANNONCE, EN_OFFRE,
                        EN_BENEFICE, EN_CLOTURE, EN_SIGNATURE])


ANGLO = {
"newsroom@thecanadianpress.com": (
    "The Canadian Press", "Fil de presse national",
    "I'm writing to the wire first because if this is worth two paragraphs, it's worth them in "
    "every paper in the country: a Quebec company insulating winter coats with the only plant "
    "monarch caterpillars can eat is going on national television Thursday.",
    "thecanadianpress.com/contact"),

"bnnassignmentdesk@bellmedia.ca": (
    "BNN Bloomberg", "Pupitre des affectations",
    "The business story: we built our own processing because no subcontractor would touch the "
    "fibre, changed our manufacturing model last year to survive, and are now pitching on "
    "national television.",
    "bnnbloomberg.ca/contact-us"),

"ROB@globeandmail.com": (
    "The Globe and Mail", "Report on Business",
    "A Quebec company built its own processing because no subcontractor would touch the fibre, "
    "changed its manufacturing model last year to survive, and is now pitching on national "
    "television. Whatever happens in the episode, the fibre and the farms behind it are real.",
    "theglobeandmail.com/about/contact"),

"globalnational@globalnews.ca": (
    "Global National", "Pupitre national",
    "A Canadian company goes on Dragons' Den this Thursday with a winter coat insulated by a "
    "weed most farmers spray — and that same weed is the only thing monarch caterpillars eat.",
    "globalnews.ca/pages/contact-us"),

"newstips@globalnews.ca": (
    "Global News", "Ligne de nouvelles",
    "A Canadian company goes on Dragons' Den this Thursday with a winter coat insulated by a "
    "weed most farmers spray — and that same weed is the only thing monarch caterpillars eat.",
    "globalnews.ca/pages/contact-us"),

"montreal@globalnews.ca": (
    "Global Montreal", "Pupitre de Montréal",
    "We're a Quebec City company, the milkweed is grown by Quebec farmers, and the insulation is "
    "made here — Thursday's national broadcast is a local story first.",
    "globalnews.ca/pages/contact-us"),

"insider@retail-insider.com": (
    "Retail Insider", "Rédaction",
    "We're a Canadian direct-to-consumer brand about to get national television exposure, with a "
    "supply chain that starts in Quebec fields — a retail story as much as an environmental one.",
    "retail-insider.com/contact-us"),

"editor@innovationintextiles.com": (
    "Innovation in Textiles", "Rédaction",
    "You covered Vegeto's milkweed insulation last February. We're the other end of that same "
    "Quebec filière: we buy the harvest, process the floss into insulation ourselves, and sell "
    "finished garments — which makes us, as far as we know, the only ones doing the whole chain.",
    "innovationintextiles.com/contact"),

"editor@knittingindustry.com": (
    "Knitting Industry", "Rédaction",
    "Milkweed floss is a hollow, naturally water-repellent fibre that behaves like down without "
    "an animal. A Quebec filière has been trying to industrialise it since 2013; we process it "
    "ourselves and sell finished garments.",
    "knittingindustry.com/contact"),

"jborneman@textileworld.com": (
    "Textile World", "Rédaction technique",
    "Milkweed floss is a hollow, naturally water-repellent fibre that behaves like down without "
    "an animal, and a Quebec filière has been trying to industrialise it since 2013. We process "
    "it ourselves and sell finished garments, which is the part nobody had managed to hold.",
    "textileworld.com/contact"),

"news@fashionunited.com": (
    "FashionUnited Canada", "Rédaction",
    "Plant-based insulation is one of the few places where outerwear can actually leave animal "
    "and petroleum fibres behind, and milkweed is the rare one that also feeds a threatened "
    "pollinator.",
    "fashionunited.ca/contact"),

"editor@betterfarming.com": (
    "Better Farming", "Rédaction",
    "The farm side: milkweed is a weed most growers spray, and a Quebec filière has been trying "
    "since 2013 to make it pay at the field gate. We still buy the harvest — that's the part that "
    "collapsed in 2018 and the part that decides whether anyone plants it again.",
    "betterfarming.com/contact"),
}


def feuille(wb, titre, entetes, lignes, largeurs):
    if titre in wb.sheetnames:
        del wb[titre]
    ws = wb.create_sheet(titre)
    ws.append(entetes)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F3864")
    for l, w in zip("ABCDEFGHIJ", largeurs):
        ws.column_dimensions[l].width = w
    for ligne in lignes:
        ws.append(ligne)
        for c in ws[ws.max_row]:
            c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[ws.max_row].height = 200
    ws.freeze_panes = "A2"
    return ws


def main(chiffrier):
    wb = openpyxl.load_workbook(chiffrier)

    fr_lignes = []
    for adresse, (nom, role) in SEMAINE_VERTE.items():
        texte = f"Bonjour {nom.split()[0]},\n\n" + fr(AGRICOLE, DISPO)
        fr_lignes.append([adresse, nom, "La Semaine verte (Radio-Canada)", role, "C",
                          "L'asclépiade à Dragons' Den le 17 sept!",
                          "ici.radio-canada.ca/tele/la-semaine-verte", None, None, texte])
    for adresse, (media, role, accroche) in TVA.items():
        texte = "Bonjour,\n\n" + fr(accroche or AFFAIRES, DISPO)
        fr_lignes.append([adresse, media, media, role, "I" if "TVA" in media else "E",
                          "L'asclépiade à Dragons' Den le 17 sept!",
                          "tvanouvelles.ca/nous-joindre", None, None, texte])
    feuille(wb, "Ajouts FR — RC et TVA",
            ["Courriel", "Nom", "Média", "Rôle", "Angle", "Objet", "Source de l'adresse",
             "Envoyé le", "Réponse", "Brouillon"],
            fr_lignes, (40, 24, 30, 24, 7, 44, 38, 12, 12, 100))

    en_lignes = []
    for adresse, (media, role, pourquoi, source) in ANGLO.items():
        en_lignes.append([adresse, media, role, "H",
                          "Quebec milkweed on Dragons' Den — this Thursday",
                          pourquoi, source, None, None, en(pourquoi)])
    feuille(wb, "Presse anglophone",
            ["Courriel", "Média", "Rôle", "Angle", "Objet", "Pourquoi eux",
             "Source de l'adresse", "Envoyé le", "Réponse", "Brouillon"],
            en_lignes, (38, 30, 26, 7, 46, 60, 34, 12, 12, 100))

    wb.save(chiffrier)
    print(f"{len(fr_lignes)} contacts FR ajoutés (La Semaine verte + TVA/Québecor), "
          f"{len(en_lignes)} anglophones, aucun @cbc.ca")


if __name__ == "__main__":
    sys.exit(main(sys.argv[1] if len(sys.argv) > 1 else "Lasclay_v2.xlsx"))
